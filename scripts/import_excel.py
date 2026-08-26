"""
Import parts and stock data from the "Stock" sheet of the legacy Stocks.xlsx ledger.

A few manual conventions used in the spreadsheet that the import must honor:
  1. Model column = "#": that slot is currently empty, no part stored there.
  2. A row's Location column is blank: that row's part is stored at the same
     physical location as the previous row.
  3. Model column = ".": that row is the same part as the previous row (just
     another batch / stored at a different location).

Usage:
  flask import-excel Stocks.xlsx           # actually import
  flask import-excel Stocks.xlsx --dry-run # stats only, no writes (rolls back at the end)
"""
import re
from collections import Counter

import click
import openpyxl

from app.extensions import db
from app.models import Location, Lot, Part, StockTransaction
from scripts.classify import build_category_lookup, classify, ensure_extra_categories

LOCATION_CODE_RE = re.compile(r"^([A-Za-z]+\d+)-(\d+)-(\d+)$")
CURRENCY_FIX = {"YSD": "USD"}


def _norm(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _get_or_create_location(code, cache):
    """Split a PREFIX+box-row-slot location code into three levels, creating/reusing nodes as needed."""
    if code in cache:
        return cache[code]

    m = LOCATION_CODE_RE.match(code)
    if not m:
        node = Location.query.filter_by(code=code, parent_id=None).first()
        if node is None:
            siblings = Location.query.filter_by(parent_id=None).count()
            node = Location(code=code, name=code, parent_id=None, sort_order=siblings)
            db.session.add(node)
            db.session.flush()
        cache[code] = node
        return node

    box, row, slot = m.group(1), m.group(2), m.group(3)
    level_codes = [box, f"{box}-{row}", f"{box}-{row}-{slot}"]

    parent = None
    node = None
    for lvl_code in level_codes:
        if lvl_code in cache:
            node = cache[lvl_code]
        else:
            parent_id = parent.id if parent else None
            node = Location.query.filter_by(code=lvl_code, parent_id=parent_id).first()
            if node is None:
                siblings = Location.query.filter_by(parent_id=parent_id).count()
                node = Location(code=lvl_code, name=lvl_code, parent_id=parent_id, sort_order=siblings)
                db.session.add(node)
                db.session.flush()
            cache[lvl_code] = node
        parent = node

    cache[code] = node
    return node


def _get_or_create_part(
    mpn, manufacturer, package, marking, description, is_fake, category_lookup, default_category_id, stats
):
    query = Part.query.filter_by(mpn=mpn)
    query = query.filter(Part.manufacturer.is_(None)) if manufacturer is None else query.filter_by(
        manufacturer=manufacturer
    )
    part = query.first()
    if part:
        stats["parts_reused"] += 1
        return part

    category_id = classify(description, category_lookup) or default_category_id
    part = Part(
        mpn=mpn,
        manufacturer=manufacturer,
        package=package,
        marking=marking,
        description=description,
        is_suspect_fake=bool(is_fake),
        category_id=category_id,
    )
    db.session.add(part)
    db.session.flush()
    stats["parts_created"] += 1
    return part


def _add_lot(part, location, qty, price, currency, stats):
    """Each Excel row that carries stock is its own purchase batch."""
    prior_total = (
        db.session.query(db.func.coalesce(db.func.sum(Lot.quantity_remaining), 0))
        .filter_by(part_id=part.id, location_id=location.id)
        .scalar()
        or 0
    )
    lot = Lot(
        part_id=part.id,
        location_id=location.id,
        quantity_received=qty,
        quantity_remaining=qty,
        unit_price=price,
        currency=currency,
    )
    db.session.add(lot)
    db.session.flush()
    db.session.add(
        StockTransaction(
            part_id=part.id,
            location_id=location.id,
            lot_id=lot.id,
            type="initial",
            quantity_delta=qty,
            balance_after=prior_total + qty,
            unit_price=price,
            currency=currency,
            reason="Imported from Excel ledger",
        )
    )
    stats["stock_rows"] += 1


def run_import(path, dry_run=False):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Stock"]

    ensure_extra_categories()
    category_lookup = build_category_lookup()
    default_category_id = category_lookup.get(("Uncategorized", None))

    location_cache = {}
    last_location_code = None
    last_part = None

    stats = Counter()

    for r in range(2, ws.max_row + 1):
        loc_raw = _norm(ws.cell(row=r, column=1).value)
        if loc_raw:
            last_location_code = loc_raw

        model_raw = ws.cell(row=r, column=2).value
        model = _norm(model_raw)
        if not model:
            continue  # fully blank separator row

        if model == "#":
            if last_location_code:
                _get_or_create_location(last_location_code, location_cache)
                stats["empty_slots"] += 1
            continue

        package = _norm(ws.cell(row=r, column=3).value)
        marking = _norm(ws.cell(row=r, column=4).value)
        manufacturer = _norm(ws.cell(row=r, column=5).value)
        qty = ws.cell(row=r, column=6).value or 0
        price = ws.cell(row=r, column=7).value
        currency = _norm(ws.cell(row=r, column=8).value) or "CNY"
        currency = CURRENCY_FIX.get(currency, currency)
        description = _norm(ws.cell(row=r, column=9).value)
        is_fake = bool(ws.cell(row=r, column=11).value)

        if model == ".":
            if last_part is None:
                stats["skipped_dot_no_ref"] += 1
                continue
            part = last_part
        else:
            part = _get_or_create_part(
                model, manufacturer, package, marking, description, is_fake,
                category_lookup, default_category_id, stats
            )
            last_part = part

        if qty and qty > 0 and last_location_code:
            location = _get_or_create_location(last_location_code, location_cache)
            _add_lot(part, location, int(qty), price, currency, stats)

    if dry_run:
        db.session.rollback()
    else:
        db.session.commit()

    click.echo(
        "Import {}: {} parts created / {} reused, {} stock batches added, {} empty slots marked, "
        "{} unresolved '.' rows, {} location nodes cached".format(
            "dry run (not written)" if dry_run else "complete",
            stats["parts_created"],
            stats["parts_reused"],
            stats["stock_rows"],
            stats["empty_slots"],
            stats["skipped_dot_no_ref"],
            len(location_cache),
        )
    )
